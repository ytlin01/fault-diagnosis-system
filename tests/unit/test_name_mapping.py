from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from services.ingestion.name_mapping import (
    apply_content_name_mapping,
    load_name_mapping,
    resolve_name_mapping_path,
)
from services.ingestion.structured_masking import read_excel_rows, run_structured_masking


class NameMappingTests(unittest.TestCase):
    def _make_temp_dir(self) -> Path:
        return Path(tempfile.mkdtemp())

    def _write_mapping_workbook(self, base_dir: Path, rows: list[list[object]]) -> Path:
        base_dir.mkdir(parents=True, exist_ok=True)
        path = base_dir / 'name_mapping.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.append(['index', 'raw_name', 'mapped_name', 'scope', 'match_mode', 'enabled', 'note', 'QPL'])
        for row in rows:
            ws.append(row)
        wb.save(path)
        wb.close()
        return path

    def _write_merged_excel(self, base_dir: Path) -> Path:
        base_dir.mkdir(parents=True, exist_ok=True)
        path = base_dir / 'merged.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = 'Sheet1'
        ws.append(['Index', 'Item', 'Description'])
        ws.append([1, 'Safety', 'Stop button pressed'])
        ws.append([2, None, 'Light curtain triggered'])
        ws.merge_cells('B2:B3')
        wb.save(path)
        wb.close()
        return path

    def test_load_mapping_and_apply_scopes_case_insensitive(self) -> None:
        temp_dir = self._make_temp_dir()
        path = self._write_mapping_workbook(temp_dir, [
            [1, 'C', 'C1STN', 'filename,sheetname', 'exact', 'TRUE', 'client alias'],
            [2, 'L', 'Vendor_A', 'content', 'exact', 'TRUE', 'vendor alias'],
            [3, 'L1', 'Line_A', 'content', 'contains', 'TRUE', 'line alias'],
            [4, 'disabled', 'ignored', 'content', 'exact', 'FALSE', 'disabled'],
        ])

        registry = load_name_mapping(path)
        self.assertEqual(registry.rule_count, 3)
        self.assertEqual(registry.scopes, ['content', 'filename', 'sheetname'])

        mapped_file, file_summary = registry.apply_text('c', scope='filename')
        self.assertEqual(mapped_file, 'C1STN')
        self.assertTrue(file_summary['applied'])

        mapped_sheet, sheet_summary = registry.apply_text('C', scope='sheetname')
        self.assertEqual(mapped_sheet, 'C1STN')
        self.assertTrue(sheet_summary['applied'])

        mapped_text, text_summary = registry.apply_text('Issue on L1', scope='content')
        self.assertEqual(mapped_text, 'Issue on Line_A')
        self.assertEqual(text_summary['replacements'], 1)

        mapped_exact, exact_summary = registry.apply_text('l', scope='content')
        self.assertEqual(mapped_exact, 'Vendor_A')
        self.assertTrue(exact_summary['applied'])

        untouched, untouched_summary = registry.apply_text('disabled', scope='content')
        self.assertEqual(untouched, 'disabled')
        self.assertFalse(untouched_summary['applied'])

    def test_apply_content_name_mapping_only_changes_configured_columns(self) -> None:
        temp_dir = self._make_temp_dir()
        path = self._write_mapping_workbook(temp_dir, [
            [1, 'L', 'Vendor_A', 'content', 'exact', 'TRUE', 'vendor alias'],
            [2, 'C', 'C1STN', 'content', 'contains', 'TRUE', 'free text alias'],
        ])
        registry = load_name_mapping(path)

        rows = [
            {
                'vendor': 'l',
                'description': 'c station alarm',
                'untouched': 'C station alarm',
            }
        ]
        mapped_rows, report = apply_content_name_mapping(
            rows,
            content_columns=['vendor', 'description'],
            registry=registry,
        )

        self.assertEqual(mapped_rows[0]['vendor'], 'Vendor_A')
        self.assertEqual(mapped_rows[0]['description'], 'C1STN station alarm')
        self.assertEqual(mapped_rows[0]['untouched'], 'C station alarm')
        self.assertEqual(report['content_rows_touched'], 1)
        self.assertEqual(report['content_cells_touched'], 2)
        self.assertEqual(report['content_replacements'], 2)

    def test_apply_content_name_mapping_adds_qpl_only_for_station_rules(self) -> None:
        temp_dir = self._make_temp_dir()
        path = self._write_mapping_workbook(temp_dir, [
            [1, 'legacy_station', 'C1STN', 'content', 'exact', 'TRUE', 'station', '1,2'],
            [2, 'acme', 'Vendor_A', 'content', 'exact', 'TRUE', 'vendor', '9'],
        ])
        registry = load_name_mapping(path)

        rows = [
            {
                'station': 'legacy_station',
                'description': 'Alarm on legacy_station',
                'vendor': 'acme',
            }
        ]
        mapped_rows, report = apply_content_name_mapping(
            rows,
            content_columns=['station', 'description', 'vendor'],
            registry=registry,
        )

        self.assertEqual(mapped_rows[0]['station'], 'C1STN')
        self.assertEqual(mapped_rows[0]['description'], 'Alarm on legacy_station')
        self.assertEqual(mapped_rows[0]['vendor'], 'Vendor_A')
        self.assertEqual(mapped_rows[0]['QPL'], [1, 2])
        self.assertEqual(report['content_rule_profile_hits']['note=station;QPL=1,2'], 1)
        self.assertEqual(report['content_rule_profile_hits']['note=vendor;QPL=9'], 1)
        self.assertEqual(report['qpl_rows_enriched'], 1)
        self.assertEqual(report['qpl_values_assigned'], 2)

    def test_read_excel_rows_expands_merged_cells(self) -> None:
        temp_dir = self._make_temp_dir()
        xlsx_path = self._write_merged_excel(temp_dir)

        raw_headers, rows, meta = read_excel_rows(
            xlsx_path,
            sheet_selector={'mode': 'index', 'value': 0},
            header_row=0,
        )

        self.assertEqual(raw_headers, ['Index', 'Item', 'Description'])
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['Item'], 'Safety')
        self.assertEqual(rows[1]['Item'], 'Safety')
        self.assertEqual(meta['sheets_read'][0]['rows'], 2)

    def test_invalid_scope_raises(self) -> None:
        temp_dir = self._make_temp_dir()
        path = self._write_mapping_workbook(temp_dir, [
            [1, 'C', 'C1STN', 'bad_scope', 'exact', 'TRUE', 'invalid scope'],
        ])
        with self.assertRaises(ValueError):
            load_name_mapping(path)

    def test_resolve_name_mapping_path_falls_back_outside_raw_dir(self) -> None:
        temp_dir = self._make_temp_dir()
        raw_dir = temp_dir / 'raw'
        raw_dir.mkdir(parents=True)
        fallback_path = self._write_mapping_workbook(temp_dir, [
            [1, 'C', 'C1STN', 'filename', 'exact', 'TRUE', 'fallback'],
        ])

        self.assertEqual(
            resolve_name_mapping_path(raw_dir),
            fallback_path.resolve(),
        )

    def test_resolve_name_mapping_path_prefers_env_var(self) -> None:
        temp_dir = self._make_temp_dir()
        raw_dir = temp_dir / 'raw'
        raw_dir.mkdir(parents=True)
        env_path = self._write_mapping_workbook(temp_dir / 'custom', [
            [1, 'C', 'C1STN', 'filename', 'exact', 'TRUE', 'env'],
        ])

        old_value = os.environ.get('FAULTDX_NAME_MAPPING_FILE')
        os.environ['FAULTDX_NAME_MAPPING_FILE'] = str(env_path)
        try:
            self.assertEqual(
                resolve_name_mapping_path(raw_dir),
                env_path,
            )
        finally:
            if old_value is None:
                os.environ.pop('FAULTDX_NAME_MAPPING_FILE', None)
            else:
                os.environ['FAULTDX_NAME_MAPPING_FILE'] = old_value

    def test_run_structured_masking_applies_mapping_before_masking(self) -> None:
        temp_dir = self._make_temp_dir()
        raw_dir = temp_dir / 'raw'
        raw_dir.mkdir(parents=True)
        output_base = temp_dir / 'out'

        csv_path = raw_dir / 'demo.csv'
        csv_path.write_text(
            'Vendor,Description\nacme,Issue on c station\n',
            encoding='utf-8',
        )

        self._write_mapping_workbook(raw_dir, [
            [1, 'DEMO.CSV', 'demo_sanitized.csv', 'filename', 'exact', 'TRUE', 'filename'],
            [2, 'ACME', 'Vendor_A', 'content', 'exact', 'TRUE', 'vendor'],
            [3, 'C', 'C1STN', 'content', 'contains', 'TRUE', 'description'],
        ])

        rules = {
            'version': 2,
            'tokenizer': {},
            'sources': {
                'demo_source': {
                    'file_patterns': ['demo.csv'],
                    'name_mapping': {
                        'content_columns': ['Vendor', 'Description'],
                    },
                    'column_rules': {
                        'Vendor': {
                            'level': 'C',
                            'canonical_action': 'passthrough',
                            'serving_action': 'passthrough',
                        },
                        'Description': {
                            'level': 'C',
                            'canonical_action': 'passthrough',
                            'serving_action': 'passthrough',
                        },
                    },
                }
            },
        }

        result = run_structured_masking(
            source_name='demo_source',
            rules=rules,
            raw_dir=raw_dir,
            output_base=output_base,
            apply_name_mapping=True,
        )

        self.assertEqual(result.canonical_rows[0]['Vendor'], 'Vendor_A')
        self.assertEqual(result.canonical_rows[0]['Description'], 'Issue on C1STN station')
        self.assertEqual(result.input_file_labels, ['demo_sanitized.csv'])
        self.assertIsNotNone(result.manifest_path)

        manifest = json.loads(result.manifest_path.read_text(encoding='utf-8'))
        self.assertEqual(manifest['run_id'], result.run_id)
        self.assertNotIn('hostname', manifest)
        lineage_record = json.loads(
            result.lineage_path.read_text(encoding='utf-8').splitlines()[0]
        )
        self.assertEqual(lineage_record['run_id'], result.run_id)
        self.assertEqual(manifest['input']['files'][0]['path'], 'demo_sanitized.csv')
        self.assertEqual(manifest['extra']['name_mapping']['content_replacements'], 2)
        self.assertEqual(
            manifest['output']['canonical']['path'],
            'out/processed/canonical/structured/demo_source/demo_source_canonical.json',
        )
        self.assertEqual(
            manifest['output']['serving']['path'],
            'out/processed/serving/structured/demo_source_serving.json',
        )

    def test_run_structured_masking_applies_governance_after_mapping(self) -> None:
        temp_dir = self._make_temp_dir()
        raw_dir = temp_dir / 'raw'
        raw_dir.mkdir(parents=True)
        output_base = temp_dir / 'out'

        csv_path = raw_dir / 'y25_demo.csv'
        csv_path.write_text(
            '序號,工站名稱,Vendor,異常描敘\n'
            '1,legacy_station,acme,Alarm on legacy_station\n'
            '2,LEGACY_STATION,acme,Alarm on legacy_station\n'
            '3,other_station,acme,Alarm on other_station\n',
            encoding='utf-8',
        )

        self._write_mapping_workbook(raw_dir, [
            [1, 'legacy_station', 'C1STN', 'content', 'exact', 'TRUE', 'station', '7,8'],
            [2, 'other_station', 'OTHER', 'content', 'exact', 'TRUE', 'station', ''],
        ])

        rules = {
            'version': 2,
            'tokenizer': {},
            'sources': {
                'y25_demo': {
                    'file_patterns': ['y25_demo.csv'],
                    'name_mapping': {
                        'content_columns': ['工站名稱', '異常描敘'],
                    },
                    'governance': {
                        'filter_equals': {
                            '工站名稱': 'C1STN',
                        },
                        'dedupe': {
                            'enabled': True,
                            'exclude_columns': ['序號'],
                        },
                    },
                    'column_rules': {
                        '序號': {
                            'level': 'C',
                            'canonical_action': 'passthrough',
                            'serving_action': 'passthrough',
                        },
                        '工站名稱': {
                            'level': 'C',
                            'canonical_action': 'passthrough',
                            'serving_action': 'passthrough',
                        },
                        'Vendor': {
                            'level': 'C',
                            'canonical_action': 'passthrough',
                            'serving_action': 'passthrough',
                        },
                        '異常描敘': {
                            'level': 'C',
                            'canonical_action': 'passthrough',
                            'serving_action': 'passthrough',
                        },
                    },
                }
            },
        }

        result = run_structured_masking(
            source_name='y25_demo',
            rules=rules,
            raw_dir=raw_dir,
            output_base=output_base,
            apply_name_mapping=True,
        )

        self.assertEqual(len(result.canonical_rows), 1)
        self.assertEqual(result.canonical_rows[0]['工站名稱'], 'C1STN')
        self.assertEqual(result.canonical_rows[0]['QPL'], [7, 8])
        self.assertEqual(result.serving_rows[0]['QPL'], [7, 8])
        self.assertEqual(result.governance_report['rows_before'], 3)
        self.assertEqual(result.governance_report['rows_after_filter'], 2)
        self.assertEqual(result.governance_report['rows_after_dedupe'], 1)
        self.assertEqual(result.governance_report['rows_filtered_out'], 1)
        self.assertEqual(result.governance_report['rows_deduped'], 1)
        self.assertIsNotNone(result.governance_history_path)

        governance_doc = json.loads(result.governance_history_path.read_text(encoding='utf-8'))
        self.assertEqual(len(governance_doc['history']['filter']['removed_rows']), 1)
        self.assertEqual(len(governance_doc['history']['dedupe']['duplicate_groups']), 1)

        manifest = json.loads(result.manifest_path.read_text(encoding='utf-8'))
        self.assertEqual(manifest['run_id'], result.run_id)
        self.assertEqual(manifest['extra']['governance']['rows_after_dedupe'], 1)
        self.assertEqual(
            manifest['extra']['name_mapping']['content_rule_profile_hits']['note=station;QPL=7,8'],
            2,
        )
        self.assertEqual(
            manifest['extra']['name_mapping']['content_rule_profile_hits']['note=station;QPL=(none)'],
            1,
        )
        self.assertEqual(manifest['extra']['name_mapping']['qpl_rows_enriched'], 2)
        self.assertEqual(manifest['extra']['name_mapping']['qpl_values_assigned'], 4)
        self.assertIn('governance_history_path', manifest['extra'])
        self.assertEqual(
            manifest['extra']['governance_history_path'],
            'out/processed/canonical/structured/y25_demo/_governance/y25_demo_governance.json',
        )


if __name__ == '__main__':
    unittest.main()
