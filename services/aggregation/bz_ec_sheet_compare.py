"""Row comparison helpers for BZ EC sheets with shared headers.

This module is intended for local structural/data review only. It compares two
Excel sheets row by row and produces a summary that can help triage whether the
two sources are aligned before handing them into fp_join-related workflows.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl


def normalize_cell(value: Any) -> str:
    """Normalize a worksheet cell into a stable comparison string."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return " ".join(text.split())


def load_sheet_rows(
    workbook_path: Path,
    *,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> dict[str, Any]:
    """Load one worksheet into headers plus row dicts.

    Args:
        workbook_path: Excel workbook path.
        sheet_name: Exact worksheet name. Defaults to the first sheet.
        header_row: 1-based header row number.
    """
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        if sheet_name is None:
            sheet_name = wb.sheetnames[0]
        if sheet_name not in wb.sheetnames:
            raise KeyError(
                f"Sheet '{sheet_name}' not found in '{workbook_path.name}'. "
                f"Available sheets: {wb.sheetnames}"
            )

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError(f"Sheet '{sheet_name}' in '{workbook_path.name}' is empty.")
        if header_row < 1 or header_row > len(rows):
            raise ValueError(
                f"header_row={header_row} is out of range for '{sheet_name}' "
                f"in '{workbook_path.name}'."
            )

        header_values = rows[header_row - 1]
        headers = [_normalize_header(value) for value in header_values]
        if not any(headers):
            raise ValueError(
                f"Header row {header_row} in '{sheet_name}' of "
                f"'{workbook_path.name}' is empty."
            )

        deduped_headers = _dedupe_headers(headers)
        records: list[dict[str, Any]] = []
        for idx, values in enumerate(rows[header_row:], start=header_row + 1):
            padded = list(values) + [None] * max(0, len(deduped_headers) - len(values))
            row_dict = {header: padded[i] if i < len(padded) else None for i, header in enumerate(deduped_headers)}
            if all(normalize_cell(value) == "" for value in row_dict.values()):
                continue
            row_dict["_row_number"] = idx
            records.append(row_dict)

        return {
            "workbook_path": str(workbook_path),
            "sheet_name": sheet_name,
            "header_row": header_row,
            "headers": deduped_headers,
            "rows": records,
        }
    finally:
        wb.close()


def compare_sheet_rows(
    left_sheet: dict[str, Any],
    right_sheet: dict[str, Any],
    *,
    key_columns: list[str] | None = None,
    ignore_columns: list[str] | None = None,
) -> dict[str, Any]:
    """Compare two loaded sheets and return a structured diff payload."""
    ignore = {column.strip() for column in (ignore_columns or []) if column.strip()}
    left_headers = [header for header in left_sheet["headers"] if header and header not in ignore]
    right_headers = [header for header in right_sheet["headers"] if header and header not in ignore]
    shared_headers = [header for header in left_headers if header in set(right_headers)]
    left_only_headers = [header for header in left_headers if header not in set(right_headers)]
    right_only_headers = [header for header in right_headers if header not in set(left_headers)]

    payload: dict[str, Any] = {
        "left": _sheet_meta(left_sheet),
        "right": _sheet_meta(right_sheet),
        "header_diff": {
            "shared_headers": shared_headers,
            "left_only_headers": left_only_headers,
            "right_only_headers": right_only_headers,
        },
    }

    if key_columns:
        missing_columns = [
            column for column in key_columns
            if column not in shared_headers
        ]
        if missing_columns:
            raise KeyError(
                "Key columns missing from shared headers: "
                + ", ".join(missing_columns)
                + ". Shared headers found: "
                + ", ".join(shared_headers)
            )
        payload["comparison"] = _compare_by_key(
            left_sheet["rows"],
            right_sheet["rows"],
            shared_headers,
            key_columns,
        )
    else:
        payload["comparison"] = _compare_by_full_row(
            left_sheet["rows"],
            right_sheet["rows"],
            shared_headers,
        )

    return payload


def _compare_by_full_row(
    left_rows: list[dict[str, Any]],
    right_rows: list[dict[str, Any]],
    columns: list[str],
) -> dict[str, Any]:
    left_counter = Counter(_row_signature(row, columns) for row in left_rows)
    right_counter = Counter(_row_signature(row, columns) for row in right_rows)

    left_only_rows: list[dict[str, Any]] = []
    right_only_rows: list[dict[str, Any]] = []

    for signature, count in (left_counter - right_counter).items():
        left_only_rows.extend(_materialize_signature(signature, columns) for _ in range(count))
    for signature, count in (right_counter - left_counter).items():
        right_only_rows.extend(_materialize_signature(signature, columns) for _ in range(count))

    return {
        "mode": "full_row",
        "row_count_summary": {
            "left_rows": len(left_rows),
            "right_rows": len(right_rows),
            "shared_exact_rows": sum((left_counter & right_counter).values()),
            "left_only_rows": len(left_only_rows),
            "right_only_rows": len(right_only_rows),
        },
        "left_only_rows": left_only_rows,
        "right_only_rows": right_only_rows,
    }


def _compare_by_key(
    left_rows: list[dict[str, Any]],
    right_rows: list[dict[str, Any]],
    shared_headers: list[str],
    key_columns: list[str],
) -> dict[str, Any]:
    compare_columns = [column for column in shared_headers if column not in set(key_columns)]
    left_index = _index_rows(left_rows, key_columns)
    right_index = _index_rows(right_rows, key_columns)

    all_keys = sorted(set(left_index) | set(right_index))
    duplicate_keys: list[dict[str, Any]] = []
    left_only_rows: list[dict[str, Any]] = []
    right_only_rows: list[dict[str, Any]] = []
    changed_rows: list[dict[str, Any]] = []
    unchanged_count = 0

    for key in all_keys:
        left_group = left_index.get(key, [])
        right_group = right_index.get(key, [])

        if len(left_group) > 1 or len(right_group) > 1:
            duplicate_keys.append(
                {
                    "key": dict(zip(key_columns, key)),
                    "left_count": len(left_group),
                    "right_count": len(right_group),
                    "left_row_numbers": [row["_row_number"] for row in left_group],
                    "right_row_numbers": [row["_row_number"] for row in right_group],
                }
            )
            continue

        if not left_group:
            right_only_rows.append(_serialize_row(right_group[0], shared_headers))
            continue
        if not right_group:
            left_only_rows.append(_serialize_row(left_group[0], shared_headers))
            continue

        left_row = left_group[0]
        right_row = right_group[0]
        diffs = []
        for column in compare_columns:
            left_value = normalize_cell(left_row.get(column))
            right_value = normalize_cell(right_row.get(column))
            if left_value != right_value:
                diffs.append(
                    {
                        "column": column,
                        "left": left_value,
                        "right": right_value,
                    }
                )
        if diffs:
            changed_rows.append(
                {
                    "key": dict(zip(key_columns, key)),
                    "left_row_number": left_row["_row_number"],
                    "right_row_number": right_row["_row_number"],
                    "diffs": diffs,
                }
            )
        else:
            unchanged_count += 1

    return {
        "mode": "keyed",
        "key_columns": key_columns,
        "compare_columns": compare_columns,
        "row_count_summary": {
            "left_rows": len(left_rows),
            "right_rows": len(right_rows),
            "unchanged_keys": unchanged_count,
            "changed_keys": len(changed_rows),
            "left_only_keys": len(left_only_rows),
            "right_only_keys": len(right_only_rows),
            "duplicate_or_ambiguous_keys": len(duplicate_keys),
        },
        "left_only_rows": left_only_rows,
        "right_only_rows": right_only_rows,
        "changed_rows": changed_rows,
        "duplicate_keys": duplicate_keys,
    }


def _index_rows(
    rows: list[dict[str, Any]],
    key_columns: list[str],
) -> dict[tuple[str, ...], list[dict[str, Any]]]:
    index: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        key = tuple(normalize_cell(row.get(column)) for column in key_columns)
        index[key].append(row)
    return index


def _row_signature(row: dict[str, Any], columns: list[str]) -> tuple[str, ...]:
    return tuple(normalize_cell(row.get(column)) for column in columns)


def _materialize_signature(signature: tuple[str, ...], columns: list[str]) -> dict[str, str]:
    return dict(zip(columns, signature))


def _serialize_row(row: dict[str, Any], columns: list[str]) -> dict[str, str]:
    data = {column: normalize_cell(row.get(column)) for column in columns}
    data["_row_number"] = str(row["_row_number"])
    return data


def _sheet_meta(sheet: dict[str, Any]) -> dict[str, Any]:
    return {
        "workbook_path": sheet["workbook_path"],
        "sheet_name": sheet["sheet_name"],
        "header_row": sheet["header_row"],
        "headers": sheet["headers"],
        "row_count": len(sheet["rows"]),
    }


def _dedupe_headers(headers: list[str]) -> list[str]:
    counts: dict[str, int] = defaultdict(int)
    deduped: list[str] = []
    for idx, header in enumerate(headers, start=1):
        base = header or f"column_{idx}"
        counts[base] += 1
        if counts[base] == 1:
            deduped.append(base)
        else:
            deduped.append(f"{base}__{counts[base]}")
    return deduped


def _normalize_header(value: Any) -> str:
    text = normalize_cell(value)
    if not text:
        return ""
    if text.startswith("*"):
        return text
    return f"*{text}"
