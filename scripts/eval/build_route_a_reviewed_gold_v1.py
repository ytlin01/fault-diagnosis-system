"""Build the V2-TASK001 attestation-reviewed gold set and scenario registry."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from services.evaluation.route_a_reviewed_gold import (
    build_reviewed_gold,
    build_scenario_registry,
)

EXPECTED_CAUSE_COUNT = 21
EXPECTED_ACTION_COUNT = 16
_CATALOG_SHEET = "Cause + Repair Review"


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _extract_catalogs(workbook_path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(workbook_path, data_only=False, read_only=True)
    if _CATALOG_SHEET not in workbook.sheetnames:
        raise SystemExit(f"workbook is missing the '{_CATALOG_SHEET}' sheet")
    sheet = workbook[_CATALOG_SHEET]
    causes: list[dict[str, Any]] = []
    actions: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=5, values_only=True):
        cause_id, cause_name, cause_definition, family = row[0], row[1], row[2], row[5]
        action_id, action_name, action_definition, action_type = (
            row[7],
            row[8],
            row[9],
            row[12],
        )
        if isinstance(cause_id, str) and cause_id.strip().startswith("cause."):
            causes.append(
                {
                    "cause_id": cause_id.strip(),
                    "display_name": str(cause_name).strip(),
                    "definition": str(cause_definition).strip() if cause_definition else "",
                    "family": str(family).strip() if family else "",
                }
            )
        if isinstance(action_id, str) and action_id.strip().startswith("action."):
            actions.append(
                {
                    "action_id": action_id.strip(),
                    "display_name": str(action_name).strip(),
                    "definition": str(action_definition).strip() if action_definition else "",
                    "action_type": str(action_type).strip() if action_type else "",
                }
            )
    workbook.close()
    if len(causes) != EXPECTED_CAUSE_COUNT:
        raise SystemExit(
            f"expected {EXPECTED_CAUSE_COUNT} cause entries, extracted {len(causes)}"
        )
    if len(actions) != EXPECTED_ACTION_COUNT:
        raise SystemExit(
            f"expected {EXPECTED_ACTION_COUNT} repair-action entries, extracted {len(actions)}"
        )
    return causes, actions


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Deterministically derive the attestation-reviewed C1STN gold set "
            "and Route A scenario registry from the annotation seed and the "
            "accepted expert-review workbook."
        )
    )
    parser.add_argument("--seed", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--gold-output", type=Path, required=True)
    parser.add_argument("--registry-output", type=Path, required=True)
    args = parser.parse_args(argv)

    seed = json.loads(args.seed.read_text(encoding="utf-8"))
    causes, actions = _extract_catalogs(args.workbook)
    seed_sha256 = _sha256_file(args.seed)
    workbook_sha256 = _sha256_file(args.workbook)

    gold = build_reviewed_gold(
        seed,
        causes,
        actions,
        seed_path=str(args.seed).replace("\\", "/"),
        seed_sha256=seed_sha256,
        workbook_path=str(args.workbook).replace("\\", "/"),
        workbook_sha256=workbook_sha256,
    )
    registry = build_scenario_registry(seed)

    for output, payload in (
        (args.gold_output, gold),
        (args.registry_output, registry),
    ):
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(f"wrote {output} sha256={_sha256_file(output)}")
    print(f"cases={len(gold['cases'])} causes={len(causes)} actions={len(actions)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
