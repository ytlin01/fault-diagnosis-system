"""Build a Schema 4.2 replacement-list candidate from the BOM04 workbook.

V2-TASK017 candidate builder. This script owns all file I/O: XLSX reading
(openpyxl), YAML config loading, salt reading, frozen-schema validation,
artifact writes, and the run manifest. The service module
``services/ingestion/material_layer_v2.py`` stays pure stdlib.

Single governed input (data-owner decision 2026-08-27): the BOM04
replacement/consumable list (L_ReplacementList). Node-only output; the
BOM03 path and topology endpoint resolution were removed.

Synthetic-only until exact governed input paths and explicit bounded data
permission are recorded in the task package. No graph import, no runtime
activation, no release claim.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))

import yaml  # noqa: E402

from services.ingestion.cad_2d_pipeline_v2.artifacts import (  # noqa: E402
    sha256_file,
    sha256_value,
    write_json_exclusive,
)
from services.ingestion import material_layer_v2 as ml  # noqa: E402

SERVICE_MODULE = PROJECT_ROOT / "services" / "ingestion" / "material_layer_v2.py"
CLI_MODULE = Path(__file__).resolve()

RUN_STATUSES = (
    "leakage_detected",
    "invalid_candidate",
    "blocked_by_exclusions",
    "candidate_with_exclusions",
    "complete_candidate",
)


class MaterialBuildError(ValueError):
    """Raised for configuration, contract, salt, validation, or output failures."""


def _load_json(path: Path) -> Any:
    try:
        with path.open("r", encoding="utf-8") as handle:
            return json.load(handle)
    except (OSError, json.JSONDecodeError) as exc:
        raise MaterialBuildError(f"cannot read JSON {path}: {exc}") from exc


def _load_workbook_tables(path: Path) -> dict[str, list[list[str | None]]]:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise MaterialBuildError(f"cannot read workbook {path}: {exc}") from exc
    try:
        tables: dict[str, list[list[str | None]]] = {}
        for sheet_name in workbook.sheetnames:
            rows: list[list[str | None]] = []
            for row in workbook[sheet_name].iter_rows(values_only=True):
                rows.append([None if cell is None else str(cell) for cell in row])
            tables[sheet_name] = rows
        return tables
    finally:
        workbook.close()


def _code_fingerprint() -> dict[str, Any]:
    """Content fingerprint of the implementation files (local freeze identity)."""
    digest = hashlib.sha256()
    files = []
    for module in (SERVICE_MODULE, CLI_MODULE):
        relative = module.relative_to(PROJECT_ROOT).as_posix()
        payload = module.read_bytes()
        digest.update(relative.encode("utf-8") + b"\0" + payload + b"\0")
        files.append(relative)
    return {"files": files, "sha256": digest.hexdigest()}


def _validate_frozen_schema(schema: Any) -> None:
    if not isinstance(schema, dict):
        raise MaterialBuildError("the Schema 4.2 document must be a JSON object")
    if schema.get("$id") != "tpdx_kg_schema_v4_2":
        raise MaterialBuildError(
            f"schema $id must be 'tpdx_kg_schema_v4_2', got {schema.get('$id')!r}"
        )
    if schema.get("x-contract-status") != "FROZEN":
        raise MaterialBuildError(
            f"schema contract status must be FROZEN, got {schema.get('x-contract-status')!r}"
        )


def _schema_validation_errors(candidate: Any, schema: dict[str, Any]) -> list[str]:
    try:
        import jsonschema
    except ModuleNotFoundError as exc:
        raise MaterialBuildError(
            "jsonschema is required to validate the candidate document"
        ) from exc
    errors = sorted(
        jsonschema.Draft202012Validator(schema).iter_errors(candidate),
        key=lambda error: [str(path) for path in error.absolute_path],
    )
    return [
        "/".join(str(part) for part in error.absolute_path) + ": " + error.message
        for error in errors
    ]


def _semantic_validation_errors(candidate: Any) -> list[str]:
    from services.governance.tpdx_schema_v4_2 import (
        SchemaV42SemanticError,
        validate_v4_2_semantics,
    )

    try:
        validate_v4_2_semantics(candidate)
    except SchemaV42SemanticError as exc:
        return [str(exc)]
    return []


def _load_salt(config: dict[str, Any]) -> str:
    salt_env = config["identity"]["salt_env"]
    salt = os.environ.get(salt_env)
    if not salt:
        raise MaterialBuildError(
            f"{salt_env} is not set. Refusing to run (fail-safe): "
            "stable material ids require the salt."
        )
    return salt


def _inspect(args: argparse.Namespace, config: dict[str, Any]) -> int:
    """Print workbook structure only: no cell values, no salt, no artifacts."""
    reports = []
    for name, path in (("bom04", args.bom04),):
        bom_cfg = config[name]
        tables = _load_workbook_tables(path)
        location = ml.locate_sheet_and_header(
            tables, bom_cfg, search_rows=config["detection"]["search_rows"]
        )
        rows = tables[location.sheet]
        headers = ml.resolve_headers(rows[location.header_row_index], bom_cfg)
        with_name = 0
        without_name = 0
        name_index = next(
            (index for index, column in headers.items() if column == bom_cfg["name_col"]), None
        )
        for row in rows[location.data_start_index:]:
            if not any(ml.canonicalize_text(cell) for cell in row):
                continue
            if name_index is not None and name_index < len(row) and ml.canonicalize_text(row[name_index]):
                with_name += 1
            else:
                without_name += 1
        reports.append({
            "input": name,
            "sheets": list(tables),
            "sheet": location.sheet,
            "fell_back": location.fell_back,
            "header_detected": location.header_detected,
            "header_row_excel": location.header_row_index + 1,
            "required_columns_present": sorted(headers.values()),
            "rows_with_name": with_name,
            "rows_without_name": without_name,
        })
    print(json.dumps({"inspect": reports}, ensure_ascii=False, indent=2))
    return 0


def _run(args: argparse.Namespace) -> int:
    try:
        with args.config.open("r", encoding="utf-8") as handle:
            raw_config = yaml.safe_load(handle)
    except (OSError, yaml.YAMLError) as exc:
        raise MaterialBuildError(f"cannot read config {args.config}: {exc}") from exc
    config = ml.validate_config(raw_config)

    if args.inspect:
        return _inspect(args, config)

    salt = _load_salt(config)
    schema = _load_json(args.schema)
    _validate_frozen_schema(schema)

    inputs_meta: dict[str, Any] = {}
    extractions: dict[str, ml.RowExtraction] = {}
    forbidden: dict[str, dict[str, Any]] = {}
    raw_names: dict[str, str] = {}
    for name, path in (("bom04", args.bom04),):
        bom_cfg = dict(config[name])
        tables = _load_workbook_tables(path)
        location = ml.locate_sheet_and_header(
            tables, bom_cfg, search_rows=config["detection"]["search_rows"]
        )
        extraction = ml.extract_rows(tables[location.sheet], bom_cfg, location)
        source_forbidden = ml.collect_forbidden_values(
            tables[location.sheet], bom_cfg, location, leakage_cfg=config["leakage"]
        )
        for value, meta in source_forbidden.items():
            forbidden.setdefault(value, meta)
        raw_names.update(extraction.raw_names)
        extractions[name] = extraction
        inputs_meta[name] = {
            "path": str(path),
            "sha256": sha256_file(path),
            "sheet": location.sheet,
            "header_row_excel": location.header_row_index + 1,
            "fell_back": location.fell_back,
        }

    result = ml.build_material_layer(
        extractions["bom04"],
        raw_names=raw_names,
        salt=salt,
        config=config,
        forbidden_values=forbidden,
    )

    json_schema_errors = _schema_validation_errors(result.candidate, schema)
    semantic_errors = _semantic_validation_errors(result.candidate)
    validation_ok = not json_schema_errors and not semantic_errors

    if result.leak_report["leak_count"]:
        run_status = "leakage_detected"
    elif not validation_ok:
        run_status = "invalid_candidate"
    elif result.exclusions and not args.allow_exclusions:
        run_status = "blocked_by_exclusions"
    elif result.exclusions:
        run_status = "candidate_with_exclusions"
    else:
        run_status = "complete_candidate"

    code = _code_fingerprint()
    result_digest = sha256_value(result.as_dict())
    identity = sha256_value({
        "config": sha256_file(args.config),
        "inputs": {name: meta["sha256"] for name, meta in sorted(inputs_meta.items())},
        "schema": sha256_file(args.schema),
        "code": code["sha256"],
        "result": result_digest,
    })
    outputs: dict[str, dict[str, Any]] = {}
    write_plan: list[tuple[Path, Any, str]] = [
        (args.audit_out, result.audit, "audit"),
        (args.lineage_out, [dict(record) for record in result.lineage], "lineage"),
        (args.exclusions_out, [dict(entry) for entry in result.exclusions], "exclusions"),
        (args.leak_report_out, result.leak_report, "leak_report"),
    ]
    if run_status in ("complete_candidate", "candidate_with_exclusions"):
        write_plan.append((args.candidate_out, result.candidate, "candidate"))

    manifest = {
        "manifest_version": "material-layer-v2-candidate-run-v1",
        "run_id": f"material_{identity[:16]}",
        "run_status": run_status,
        "promotion_state": "candidate",
        "graph_activation": False,
        "config": {
            "path": str(args.config),
            "sha256": sha256_file(args.config),
            "schema_version": config["schema_version"],
        },
        "inputs": inputs_meta,
        "schema_v4_2": {
            "path": str(args.schema),
            "sha256": sha256_file(args.schema),
            "contract_status": "FROZEN",
        },
        "code": code,
        "salt_source": config["identity"]["salt_env"],
        "identity_rule_version": config["identity"]["identity_rule_version"],
        "outputs": outputs,
        "validation": {
            "json_schema": "passed" if not json_schema_errors else "failed",
            "json_schema_errors": json_schema_errors,
            "semantic": "passed" if not semantic_errors else "failed",
            "semantic_errors": semantic_errors,
            "forbidden_values_checked": result.leak_report["forbidden_values_checked"],
            "leak_count": result.leak_report["leak_count"],
        },
        "counts": result.audit["counts"],
        "rollback": {"graph_import": False, "activation_required": False},
    }

    written: list[Path] = []
    try:
        for path, value, label in write_plan:
            write_json_exclusive(path, value)
            written.append(path)
            outputs[label] = {
                "path": str(path),
                "sha256": sha256_file(path),
                "written": True,
            }
        outputs["candidate"] = outputs.get(
            "candidate",
            {"path": str(args.candidate_out), "sha256": None, "written": False},
        )
        write_json_exclusive(args.manifest_out, manifest)
        written.append(args.manifest_out)
    except FileExistsError as exc:
        for path in reversed(written):
            path.unlink(missing_ok=True)
        raise MaterialBuildError(
            f"refusing to overwrite existing artifact ({exc}); rerun into a fresh output directory"
        ) from exc
    except Exception:
        for path in reversed(written):
            path.unlink(missing_ok=True)
        raise

    print(json.dumps(manifest, ensure_ascii=False, sort_keys=True))
    if run_status == "blocked_by_exclusions":
        print(
            "blocked by governed exclusions; rerun with --allow-exclusions to write the "
            "review candidate",
            file=sys.stderr,
        )
        return 3
    if run_status in ("leakage_detected", "invalid_candidate"):
        print(f"ERROR: run status {run_status}; evidence written, candidate withheld", file=sys.stderr)
        return 2
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Build a Schema 4.2 material-layer candidate (V2-TASK017, synthetic-only)."
    )
    parser.add_argument("--config", required=True, type=Path)
    parser.add_argument("--bom04", required=True, type=Path, help="BOM04 replacement-list workbook (.xlsx)")
    parser.add_argument("--schema", type=Path, help="frozen Schema 4.2 JSON schema")
    parser.add_argument("--candidate-out", type=Path)
    parser.add_argument("--exclusions-out", type=Path)
    parser.add_argument("--lineage-out", type=Path)
    parser.add_argument("--audit-out", type=Path)
    parser.add_argument("--leak-report-out", type=Path)
    parser.add_argument("--manifest-out", type=Path)
    parser.add_argument("--allow-exclusions", action="store_true")
    parser.add_argument("--inspect", action="store_true")
    args = parser.parse_args(argv)

    if not args.inspect:
        missing = [
            name
            for name in (
                "candidate_out", "exclusions_out", "lineage_out", "audit_out",
                "leak_report_out", "manifest_out",
            )
            if getattr(args, name) is None
        ]
        if missing:
            parser.error(f"--inspect not set; missing required arguments: {', '.join(missing)}")

    try:
        return _run(args)
    except MaterialBuildError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
