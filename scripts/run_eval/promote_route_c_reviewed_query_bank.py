"""Convert the reviewed Route C query-bank workbook into an eval JSON bank."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DEFAULT_INPUT = Path("data/eval/route_c/route_c_human_query_bank_review_v1.xlsx")
DEFAULT_OUTPUT = Path("data/eval/route_c/route_c_human_query_bank_reviewed_v1.json")


def promote(input_path: Path, output_path: Path) -> dict[str, Any]:
    worksheet = load_workbook(input_path, read_only=True, data_only=True)["query_bank_review"]
    rows = list(worksheet.iter_rows(values_only=True))
    headers = [str(value) for value in rows[0]]
    cases = [dict(zip(headers, row)) for row in rows[1:] if any(value not in (None, "") for value in row)]
    if any(str(case.get("gold_status") or "").strip().lower() != "reviewed" for case in cases):
        raise ValueError("All query cases must have gold_status=reviewed.")
    normalized = []
    for case in cases:
        normalized.append({
            "query_id": case["query_id"], "query": case["query"], "query_type": case["query_type"],
            "difficulty": case["difficulty"], "query_language": case["query_language"],
            "expected_rag_candidate_ids": _split(case.get("expected_rag_candidate_ids")),
            "acceptable_root_causes": _split(case.get("acceptable_root_causes")),
            "acceptable_countermeasures": _split(case.get("acceptable_countermeasures")),
            "source_rag_candidate_id": case.get("source_rag_candidate_id"),
            "review_status": "reviewed", "gold_status": "reviewed", "review_notes": case.get("review_notes") or "",
        })
    bank = {"schema_version": "route_c_human_query_bank_reviewed_v1", "route": "route_c", "source_workbook": input_path.as_posix(), "status": "reviewed", "gold_status": "reviewed", "cases": normalized}
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(bank, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return bank


def _split(value: Any) -> list[str]:
    return [part.strip() for part in str(value or "").split(" || ") if part.strip()]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    bank = promote(args.input, args.output)
    print(json.dumps({"status": bank["status"], "cases": len(bank["cases"]), "output": args.output.as_posix()}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
