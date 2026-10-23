"""Build the human-facing V2-TASK001 C1STN domain-expert review workbook.

This is a presentation and annotation aid for the governed JSON seed.  It does
not create canonical labels or promote source-suggested repair actions.
"""

from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


NAVY = "17365D"
BLUE = "D9EAF7"
YELLOW = "FFF2CC"
ORANGE = "FCE4D6"
GREEN = "E2F0D9"
GREY = "F2F2F2"
WHITE = "FFFFFF"
LIGHT_BORDER = Side(style="thin", color="D9E2F3")


def style_title(ws, cell_range: str, title: str) -> None:
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = title
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True, size=16)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def style_header(ws, row: int, start_col: int, end_col: int) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=LIGHT_BORDER)


def set_widths(ws, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def build_workbook(seed_path: Path, output_path: Path) -> None:
    with seed_path.open(encoding="utf-8") as seed_file:
        seed = json.load(seed_file)
    cases = seed["cases"]

    workbook = Workbook()
    start = workbook.active
    start.title = "Start Here"
    start.sheet_view.showGridLines = False
    style_title(start, "A1:H2", "V2-TASK001 | C1STN Domain Expert Review Workbook")
    start["A4"] = "Purpose"
    start["A4"].font = Font(bold=True, color=NAVY)
    start["B4"] = (
        "Review masked source cases independently. This workbook captures proposed labels; "
        "it does not contain approved gold labels."
    )
    start["A6"] = "Important boundary"
    start["A6"].font = Font(bold=True, color=NAVY)
    start["B6"] = (
        "The source-suggested repair action is reference-only. Do not treat it as diagnostic truth. "
        "Do not copy this workbook into a serving system or knowledge graph."
    )
    start["A8"] = "How to review"
    start["A8"].font = Font(bold=True, color=NAVY)
    instructions = [
        "1. Open Review Queue and filter for your assigned cases.",
        "2. Read the masked error context and the review query.",
        "3. Enter a proposed cause and repair-action ID/name, or leave blank and mark Needs information.",
        "4. Record confidence, rationale, reviewer ID, and review date.",
        "5. Add any new proposed IDs to Vocabulary Proposals. They are not canonical until approved.",
        "6. Return the workbook to the governed annotation owner for comparison and adjudication.",
    ]
    for index, instruction in enumerate(instructions, start=9):
        start.cell(row=index, column=2, value=instruction)
    start["A16"] = "Editable cells"
    start["A16"].font = Font(bold=True, color=NAVY)
    start["B16"] = "Yellow cells are expert inputs. Orange cells are source-suggested reference only."
    start["A18"] = "Seed"
    start["B18"] = seed["queue_id"]
    start["A19"] = "Cases"
    start["B19"] = len(cases)
    start["A20"] = "Seed state"
    start["B20"] = seed["queue_state"]
    for row in range(4, 21):
        start.cell(row=row, column=2).alignment = Alignment(vertical="top", wrap_text=True)
    start["B6"].fill = PatternFill("solid", fgColor=ORANGE)
    start["B16"].fill = PatternFill("solid", fgColor=YELLOW)
    set_widths(start, {"A": 23, "B": 105, "C": 3, "D": 18, "E": 18, "F": 18, "G": 18, "H": 18})
    for row in (4, 6, 8, 16):
        start.row_dimensions[row].height = 30

    queue = workbook.create_sheet("Review Queue")
    queue.sheet_view.showGridLines = False
    style_title(queue, "A1:T2", "C1STN Error-Code Review Queue | Source Context Is Reference Only")
    queue["A3"] = "Orange columns: masked source context / source suggestion. Yellow columns: reviewer input."
    queue.merge_cells("A3:T3")
    queue["A3"].fill = PatternFill("solid", fgColor=GREY)
    queue["A3"].font = Font(italic=True, color="595959")

    headers = [
        "Case ID", "Scenario ID", "Error Code", "Local Alarm Code", "Category",
        "Error Description", "Error Description (Chinese)", "Source-Suggested Repair Action (reference only)",
        "Review Query", "Difficulty", "Proposed Cause ID", "Proposed Cause Name",
        "Proposed Repair Action ID", "Proposed Repair Action Name", "Decision",
        "Confidence", "Rationale", "Reviewer ID", "Review Date", "Review Status",
    ]
    for col, header in enumerate(headers, start=1):
        queue.cell(row=5, column=col, value=header)
    style_header(queue, 5, 1, len(headers))

    for row_index, case in enumerate(cases, start=6):
        context = case["source_context"]
        values = [
            case["case_id"], case["scenario_id"], context["code"], context["local_alarm_code"], context["category"],
            context["error_description"], context["error_description_zh"], context["source_suggested_repair_action"],
            case["query"], case["coverage"]["difficulty"], "", "", "", "", "", "", "", "", "", "Not started",
        ]
        for col, value in enumerate(values, start=1):
            cell = queue.cell(row=row_index, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=LIGHT_BORDER)
        for col in range(3, 11):
            queue.cell(row=row_index, column=col).fill = PatternFill("solid", fgColor=ORANGE)
        for col in range(11, 21):
            queue.cell(row=row_index, column=col).fill = PatternFill("solid", fgColor=YELLOW)
        queue.cell(row=row_index, column=19).number_format = "yyyy-mm-dd"
        queue.row_dimensions[row_index].height = 70

    last_row = len(cases) + 5
    queue.auto_filter.ref = f"A5:T{last_row}"
    queue.freeze_panes = "C6"
    set_widths(queue, {
        "A": 22, "B": 22, "C": 14, "D": 17, "E": 20, "F": 34, "G": 34, "H": 38,
        "I": 42, "J": 13, "K": 27, "L": 30, "M": 31, "N": 32, "O": 20, "P": 15,
        "Q": 40, "R": 20, "S": 14, "T": 18,
    })

    for column, choices in {
        "O": ["Accept proposal", "Reject proposal", "Needs information"],
        "P": ["High", "Medium", "Low"],
        "T": ["Not started", "In progress", "Completed", "Escalate"],
    }.items():
        validation = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        validation.error = "Select a value from the list."
        validation.errorTitle = "Invalid review value"
        queue.add_data_validation(validation)
        validation.add(f"{column}6:{column}{last_row}")

    queue.conditional_formatting.add(
        f"T6:T{last_row}",
        FormulaRule(formula=["$T6=\"Completed\""], fill=PatternFill("solid", fgColor=GREEN)),
    )

    guide = workbook.create_sheet("ID Proposal Guide")
    guide.sheet_view.showGridLines = False
    style_title(guide, "A1:F2", "Proposed ID Guidance | Not a Canonical Vocabulary")
    guide["A4"] = "Use these fields only if no approved controlled vocabulary has been supplied."
    guide.merge_cells("A4:F4")
    guide["A4"].fill = PatternFill("solid", fgColor=ORANGE)
    guide["A4"].alignment = Alignment(wrap_text=True)
    guide_headers = ["Proposal type", "Suggested format", "Definition to provide", "Do not do", "Approval status", "Example format only"]
    for col, header in enumerate(guide_headers, start=1):
        guide.cell(row=6, column=col, value=header)
    style_header(guide, 6, 1, 6)
    guide_rows = [
        ["Cause", "cause.c1stn.<short_slug>", "A concise, mechanism-focused failure cause.", "Do not use an error code alone as the cause definition.", "Proposed until approved", "cause.c1stn.link_communication_loss"],
        ["Repair action", "action.c1stn.<short_slug>", "A concrete, safely scoped corrective action.", "Do not copy the source suggestion without independent judgment.", "Proposed until approved", "action.c1stn.inspect_link_connection"],
    ]
    for row, values in enumerate(guide_rows, start=7):
        for col, value in enumerate(values, start=1):
            cell = guide.cell(row=row, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=LIGHT_BORDER)
    guide["A11"] = "Approval path"
    guide["A11"].font = Font(bold=True, color=NAVY)
    guide["B11"] = "Expert proposal -> canonicalization owner review -> controlled-vocabulary approval -> annotation adjudication"
    guide["B11"].alignment = Alignment(wrap_text=True)
    set_widths(guide, {"A": 18, "B": 34, "C": 45, "D": 48, "E": 27, "F": 42})
    for row in (4, 7, 8, 11):
        guide.row_dimensions[row].height = 42

    proposals = workbook.create_sheet("Vocabulary Proposals")
    proposals.sheet_view.showGridLines = False
    style_title(proposals, "A1:H2", "Vocabulary Proposals | Pending Canonicalization Review")
    proposal_headers = ["Proposal ID", "Proposal type", "Preferred name", "Definition", "Evidence case ID", "Proposed by", "Review status", "Canonicalization notes"]
    for col, header in enumerate(proposal_headers, start=1):
        proposals.cell(row=4, column=col, value=header)
    style_header(proposals, 4, 1, len(proposal_headers))
    for row in range(5, 25):
        for col in range(1, 9):
            proposals.cell(row=row, column=col).fill = PatternFill("solid", fgColor=YELLOW)
            proposals.cell(row=row, column=col).alignment = Alignment(vertical="top", wrap_text=True)
            proposals.cell(row=row, column=col).border = Border(bottom=LIGHT_BORDER)
        proposals.row_dimensions[row].height = 36
    proposal_validation = DataValidation(type="list", formula1='"Cause,Repair action"', allow_blank=True)
    proposals.add_data_validation(proposal_validation)
    proposal_validation.add("B5:B24")
    status_validation = DataValidation(type="list", formula1='"Proposed,Approved,Rejected,Needs revision"', allow_blank=True)
    proposals.add_data_validation(status_validation)
    status_validation.add("G5:G24")
    set_widths(proposals, {"A": 35, "B": 18, "C": 30, "D": 55, "E": 24, "F": 22, "G": 22, "H": 48})
    proposals.freeze_panes = "A5"

    for sheet in workbook.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.4
        sheet.page_margins.bottom = 0.4

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def verify_workbook(output_path: Path, expected_cases: int) -> None:
    workbook = load_workbook(output_path, data_only=False)
    expected_sheets = ["Start Here", "Review Queue", "ID Proposal Guide", "Vocabulary Proposals"]
    if workbook.sheetnames != expected_sheets:
        raise ValueError(f"Unexpected sheets: {workbook.sheetnames}")
    queue = workbook["Review Queue"]
    if queue.max_row != expected_cases + 5:
        raise ValueError(f"Expected {expected_cases} review rows; found {queue.max_row - 5}")
    if queue["K6"].value or queue["M6"].value:
        raise ValueError("Workbook must not pre-populate proposed IDs")
    if queue["H6"].value is None:
        raise ValueError("Missing source-suggested reference context")
    if workbook["Vocabulary Proposals"]["A5"].value is not None:
        raise ValueError("Vocabulary proposal table must start blank")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--seed",
        default="data/processed/review_pending/route_a/v2_task001_c1stn_error_code_annotation_seed_v1.json",
    )
    parser.add_argument(
        "--output",
        default="docs/outputs/v2_task001_c1stn_domain_review_20260723/V2-TASK001_C1STN_Domain_Expert_Review.xlsx",
    )
    args = parser.parse_args()
    seed_path = Path(args.seed)
    output_path = Path(args.output)
    with seed_path.open(encoding="utf-8") as seed_file:
        expected_cases = len(json.load(seed_file)["cases"])
    build_workbook(seed_path, output_path)
    verify_workbook(output_path, expected_cases)
    print(f"Created and verified: {output_path}")


if __name__ == "__main__":
    main()
