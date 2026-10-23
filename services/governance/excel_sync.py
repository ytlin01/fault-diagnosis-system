"""XLSX review-export formatting helper for canonical governance records.

Maintains five sheets:
    source_inventory     — per-artifact registration (zone / masking / eligibility)
    masking_log          — per-run masking history
    annotation_log       — reviewer identity + source version binding
    index_registry       — index build ↔ serving snapshot binding
    eval_run_log         — eval run records

All mutations are file-level (open → modify → save) so no persistent connection
is needed. Not thread-safe — single-process CLI use only.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from services.governance.schema import SHEET_HEADERS

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


# ── ExcelSync class ───────────────────────────────────────────────────────────

class ExcelSync:
    """Write and read a generated XLSX review export; never use it as canonical storage."""

    def __init__(self, path: Path) -> None:
        self.path = path

    # ── Initialization ────────────────────────────────────────────────────────

    def initialize(self, *, overwrite: bool = False) -> None:
        """Create the workbook with all five sheets and headers.

        If the file already exists, does nothing unless overwrite=True.
        """
        if self.path.exists() and not overwrite:
            return

        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()

        # Remove default sheet
        default = wb.active
        if default is not None:
            wb.remove(default)

        for sheet_name, headers in SHEET_HEADERS.items():
            ws = wb.create_sheet(sheet_name)
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = _HEADER_FONT
                cell.fill = _HEADER_FILL
            # Freeze header row
            ws.freeze_panes = "A2"

        wb.save(self.path)
        print(f"[governance] Initialized {self.path} with {len(SHEET_HEADERS)} sheets.")

    # ── Write ─────────────────────────────────────────────────────────────────

    def append_row(self, sheet_name: str, row: dict[str, Any]) -> None:
        """Append one row dict to the given sheet.

        Keys in row must match the sheet's header columns (extra keys ignored,
        missing keys written as empty string).
        """
        if sheet_name not in SHEET_HEADERS:
            raise ValueError(f"Unknown sheet: '{sheet_name}'. Valid: {list(SHEET_HEADERS)}")

        wb = openpyxl.load_workbook(self.path)
        ws = wb[sheet_name]
        headers = SHEET_HEADERS[sheet_name]
        values = ["" if row.get(h) is None else str(row.get(h, "")) for h in headers]
        ws.append(values)
        last_col = get_column_letter(len(headers))
        for table in ws.tables.values():
            table.ref = f"A1:{last_col}{ws.max_row}"
        wb.save(self.path)

    def update_row(
        self,
        sheet_name: str,
        key_col: str,
        key_val: str,
        updates: dict[str, Any],
    ) -> bool:
        """Update fields on the first row where key_col == key_val.

        Returns True if a row was found and updated, False otherwise.
        """
        if sheet_name not in SHEET_HEADERS:
            raise ValueError(f"Unknown sheet: '{sheet_name}'")

        headers = SHEET_HEADERS[sheet_name]
        if key_col not in headers:
            raise ValueError(f"Column '{key_col}' not in sheet '{sheet_name}'")

        key_idx = headers.index(key_col)
        wb = openpyxl.load_workbook(self.path)
        ws = wb[sheet_name]

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
            if str(row[key_idx].value or "") == str(key_val):
                for col_name, new_val in updates.items():
                    if col_name in headers:
                        col_idx = headers.index(col_name) + 1
                        ws.cell(row=row_idx, column=col_idx, value=str(new_val or ""))
                wb.save(self.path)
                return True

        return False

    # ── Read ──────────────────────────────────────────────────────────────────

    def read_sheet(self, sheet_name: str) -> list[dict[str, str]]:
        """Return all data rows from a sheet as a list of dicts.

        Uses the first row as headers.
        """
        if sheet_name not in SHEET_HEADERS:
            raise ValueError(f"Unknown sheet: '{sheet_name}'")

        wb = openpyxl.load_workbook(self.path, data_only=True, read_only=True)
        try:
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            header_row = next(rows_iter, None)
            if header_row is None:
                return []
            headers = [str(h or "") for h in header_row]
            result: list[dict[str, str]] = []
            for row in rows_iter:
                if all(v is None for v in row):
                    continue
                result.append({
                    headers[i]: str(row[i] or "")
                    for i in range(min(len(headers), len(row)))
                })
            return result
        finally:
            wb.close()

    def row_exists(self, sheet_name: str, key_col: str, key_val: str) -> bool:
        """Return True if any row in sheet_name has key_col == key_val."""
        for row in self.read_sheet(sheet_name):
            if row.get(key_col, "") == str(key_val):
                return True
        return False

    # ── Timestamp helper ──────────────────────────────────────────────────────

    @staticmethod
    def now_iso() -> str:
        return datetime.now(timezone.utc).strftime("%Y-%m-%d")

    @staticmethod
    def now_timestamp() -> str:
        return datetime.now(timezone.utc).isoformat(timespec="seconds")
