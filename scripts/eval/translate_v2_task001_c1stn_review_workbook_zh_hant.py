"""Create a Traditional-Chinese expert-review copy of the V2-TASK001 workbook."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


def replace_validation(sheet, cell_range: str, choices: list[str]) -> None:
    validation = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
    validation.error = "請從清單中選取一個值。"
    validation.errorTitle = "無效的審查值"
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def build_translation(source_path: Path, output_path: Path) -> None:
    workbook = load_workbook(source_path)
    workbook["Start Here"].title = "開始使用"
    workbook["Review Queue"].title = "審查清單"
    workbook["ID Proposal Guide"].title = "ID 提案指引"
    workbook["Vocabulary Proposals"].title = "詞彙提案"

    start = workbook["開始使用"]
    start["A1"] = "V2-TASK001｜C1STN 領域專家審查工作簿"
    translations = {
        "A4": "目的",
        "B4": "獨立審查已遮罩的來源案例。本工作簿用於記錄擬議標籤；當中不含已核准的黃金標籤。",
        "A6": "重要界線",
        "B6": "來源建議的修復動作僅供參考，請勿視為診斷事實。請勿將本工作簿複製到服務系統或知識圖譜。",
        "A8": "審查步驟",
        "B9": "1. 開啟「審查清單」，並篩選您被分配的案例。",
        "B10": "2. 閱讀已遮罩的錯誤情境與審查問題。",
        "B11": "3. 填寫擬議的原因與修復動作 ID／名稱；若資訊不足，請留白並標記「需要更多資訊」。",
        "B12": "4. 記錄信心程度、理由、審查者 ID 與審查日期。",
        "B13": "5. 將新的擬議 ID 加入「詞彙提案」。在核准前，它們不是正式標準詞彙。",
        "B14": "6. 將完成的工作簿交回受管治理的標註負責人，以便比對及裁決。",
        "A16": "可編輯儲存格",
        "B16": "黃色儲存格供專家填寫；橙色儲存格為來源建議，僅供參考。",
        "A18": "種子佇列",
        "A19": "案例數",
        "A20": "種子狀態",
    }
    for cell, value in translations.items():
        start[cell] = value

    queue = workbook["審查清單"]
    queue["A1"] = "C1STN 錯誤代碼審查清單｜來源情境僅供參考"
    queue["A3"] = "橙色欄位：已遮罩的來源情境／來源建議。黃色欄位：專家填寫。"
    headers = [
        "案例 ID", "情境 ID", "錯誤代碼", "本地警報代碼", "類別", "錯誤說明",
        "錯誤說明（中文）", "來源建議的修復動作（僅供參考）", "審查問題", "難度",
        "擬議原因 ID", "擬議原因名稱", "擬議修復動作 ID", "擬議修復動作名稱",
        "決定", "信心程度", "理由", "審查者 ID", "審查日期", "審查狀態",
    ]
    for column, header in enumerate(headers, start=1):
        queue.cell(row=5, column=column, value=header)
    difficulty_map = {"high": "高", "medium": "中", "low": "低"}
    status_map = {"Not started": "尚未開始", "In progress": "進行中", "Completed": "已完成", "Escalate": "需升級處理"}
    for row in range(6, queue.max_row + 1):
        queue.cell(row=row, column=10).value = difficulty_map.get(queue.cell(row=row, column=10).value, queue.cell(row=row, column=10).value)
        queue.cell(row=row, column=20).value = status_map.get(queue.cell(row=row, column=20).value, queue.cell(row=row, column=20).value)
    queue.data_validations.dataValidation = []
    replace_validation(queue, f"O6:O{queue.max_row}", ["接受提案", "拒絕提案", "需要更多資訊"])
    replace_validation(queue, f"P6:P{queue.max_row}", ["高", "中", "低"])
    replace_validation(queue, f"T6:T{queue.max_row}", ["尚未開始", "進行中", "已完成", "需升級處理"])
    queue.conditional_formatting._cf_rules.clear()
    queue.conditional_formatting.add(
        f"T6:T{queue.max_row}",
        FormulaRule(formula=['$T6="已完成"'], fill=PatternFill("solid", fgColor="E2F0D9")),
    )

    guide = workbook["ID 提案指引"]
    guide["A1"] = "擬議 ID 指引｜這不是正式的受控詞彙表"
    guide["A4"] = "若尚未提供已核准的受控詞彙表，才可使用這些欄位。"
    guide_headers = ["提案類型", "建議格式", "應提供的定義", "請勿這樣做", "核准狀態", "僅示範格式"]
    for column, header in enumerate(guide_headers, start=1):
        guide.cell(row=6, column=column, value=header)
    guide_rows = [
        ["原因", "cause.c1stn.<短代稱>", "簡潔、以失效機制為中心的原因定義。", "請勿只用錯誤代碼作為原因定義。", "核准前均為擬議", "cause.c1stn.link_communication_loss"],
        ["修復動作", "action.c1stn.<短代稱>", "具體且安全範圍明確的修復動作。", "請勿未經獨立判斷便直接採用來源建議。", "核准前均為擬議", "action.c1stn.inspect_link_connection"],
    ]
    for row, values in enumerate(guide_rows, start=7):
        for column, value in enumerate(values, start=1):
            guide.cell(row=row, column=column, value=value)
    guide["A11"] = "核准流程"
    guide["B11"] = "專家提案 → 正規化負責人審查 → 受控詞彙表核准 → 標註裁決"

    proposals = workbook["詞彙提案"]
    proposals["A1"] = "詞彙提案｜等待正規化審查"
    proposal_headers = ["提案 ID", "提案類型", "建議名稱", "定義", "佐證案例 ID", "提案者", "審查狀態", "正規化備註"]
    for column, header in enumerate(proposal_headers, start=1):
        proposals.cell(row=4, column=column, value=header)
    proposals.data_validations.dataValidation = []
    replace_validation(proposals, "B5:B24", ["原因", "修復動作"])
    replace_validation(proposals, "G5:G24", ["擬議", "已核准", "已拒絕", "需要修訂"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def verify_translation(output_path: Path) -> None:
    workbook = load_workbook(output_path, data_only=False)
    expected_sheets = ["開始使用", "審查清單", "ID 提案指引", "詞彙提案"]
    if workbook.sheetnames != expected_sheets:
        raise ValueError(f"Unexpected sheets: {workbook.sheetnames}")
    queue = workbook["審查清單"]
    if queue.max_row != 20 or queue.max_column != 20:
        raise ValueError("Unexpected queue dimensions")
    if queue["K6"].value or queue["M6"].value:
        raise ValueError("Translated workbook must not pre-populate proposed IDs")
    if queue["T6"].value != "尚未開始":
        raise ValueError("Review status was not translated")
    if len(queue.data_validations.dataValidation) != 3:
        raise ValueError("Expected three review controls")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        default="docs/outputs/v2_task001_c1stn_domain_review_20260723/V2-TASK001_C1STN_Domain_Expert_Review.xlsx",
    )
    parser.add_argument(
        "--output",
        default="docs/outputs/v2_task001_c1stn_domain_review_20260723/V2-TASK001_C1STN_領域專家審查工作簿_繁體中文.xlsx",
    )
    args = parser.parse_args()
    build_translation(Path(args.source), Path(args.output))
    verify_translation(Path(args.output))
    print(f"Created and verified: {args.output}")


if __name__ == "__main__":
    main()
